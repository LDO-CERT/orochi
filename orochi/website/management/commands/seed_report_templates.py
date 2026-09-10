import io

from django.core.files.base import ContentFile
from django.core.management.base import BaseCommand
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

from orochi.website.models import ReportTemplate


def set_cell_background(cell, hex_color):
    """Set background color of a docx table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner cell padding."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f"<w:tcMar {nsdecls('w')}>"
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f"</w:tcMar>"
    )
    tc_pr.append(tc_mar)


def create_default_docx_bytes():
    doc = Document()

    # Set normal style font
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Document Header / Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    t_run = title.add_run("OROCHI FORENSIC INCIDENT REPORT")
    t_run.font.size = Pt(22)
    t_run.font.bold = True
    t_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)  # Navy blue

    sub = doc.add_paragraph()
    s_run = sub.add_run("Case: {{ case.name }}")
    s_run.font.size = Pt(13)
    s_run.font.bold = True
    s_run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 1: Case Overview
    h1 = doc.add_heading("Case Overview", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    ov_table = doc.add_table(rows=5, cols=2)
    ov_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    fields = [
        ("Case Name", "{{ case.name }}"),
        ("Status", "{{ case.status }}"),
        ("Lead Investigator", "{{ case.user.username }}"),
        ("Created Date", "{{ case.created_at }}"),
        ("Description", '{{ case.description or "-" }}'),
    ]
    for idx, (label, val) in enumerate(fields):
        row = ov_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.0)
        c1.width = Inches(4.5)
        set_cell_background(c0, "F3F4F6")
        set_cell_margins(c0)
        set_cell_margins(c1)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        p1 = c1.paragraphs[0]
        p1.add_run(val)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 2: Executive Summary
    h2 = doc.add_heading("Executive Summary", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_ai = doc.add_paragraph()
    p_ai.paragraph_format.line_spacing = 1.15
    p_ai.paragraph_format.space_after = Pt(12)
    r_ai = p_ai.add_run('{{ ai_summary or "No executive summary has been generated for this case." }}')
    r_ai.font.italic = False

    # Section 3: Evidence Items
    h3 = doc.add_heading("Attached Evidence", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)
    for r in h3.runs:
        r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    ev_table = doc.add_table(rows=4, cols=3)
    ev_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header row
    hdr_cells = ev_table.rows[0].cells
    hdr_cells[0].width = Inches(2.2)
    hdr_cells[1].width = Inches(1.8)
    hdr_cells[2].width = Inches(2.5)
    for c, text in zip(hdr_cells, ["Evidence Name", "Memory Dump", "Description"], strict=False):
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, top=120, bottom=120)
        p = c.paragraphs[0]
        run = p.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Loop row start
    r1 = ev_table.rows[1]
    r1.cells[0].text = "{%tr for e in evidences %}"
    # Data row
    r2 = ev_table.rows[2]
    r2.cells[0].width = Inches(2.2)
    r2.cells[1].width = Inches(1.8)
    r2.cells[2].width = Inches(2.5)
    r2.cells[0].text = "{{ e.name }}"
    r2.cells[1].text = '{{ e.dump.name if e.dump else "-" }}'
    r2.cells[2].text = '{{ e.description or "-" }}'
    for c in r2.cells:
        set_cell_margins(c)
    # Loop row end
    r3 = ev_table.rows[3]
    r3.cells[0].text = "{%tr endfor %}"

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 4: Forensic Findings
    h4 = doc.add_heading("Forensic Findings", level=1)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)
    for r in h4.runs:
        r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    f_table = doc.add_table(rows=4, cols=3)
    f_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    f_hdr = f_table.rows[0].cells
    f_hdr[0].width = Inches(1.5)
    f_hdr[1].width = Inches(1.8)
    f_hdr[2].width = Inches(3.2)
    for c, text in zip(f_hdr, ["Severity", "MITRE ATT&CK", "Observations & Notes"], strict=False):
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, top=120, bottom=120)
        p = c.paragraphs[0]
        run = p.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    f_table.rows[1].cells[0].text = "{%tr for f in findings %}"
    fr2 = f_table.rows[2]
    fr2.cells[0].width = Inches(1.5)
    fr2.cells[1].width = Inches(1.8)
    fr2.cells[2].width = Inches(3.2)
    fr2.cells[0].text = "{{ f.severity }}"
    fr2.cells[1].text = '{{ f.mitre_attack_technique or "-" }}'
    fr2.cells[2].text = "{{ f.note }}"
    for c in fr2.cells:
        set_cell_margins(c)
    f_table.rows[3].cells[0].text = "{%tr endfor %}"

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 5: Incident Timeline
    h5 = doc.add_heading("Incident Timeline", level=1)
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)
    for r in h5.runs:
        r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    t_table = doc.add_table(rows=4, cols=3)
    t_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_hdr = t_table.rows[0].cells
    t_hdr[0].width = Inches(2.2)
    t_hdr[1].width = Inches(1.8)
    t_hdr[2].width = Inches(2.5)
    for c, text in zip(t_hdr, ["Timestamp", "Event Type", "Event Description"], strict=False):
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, top=120, bottom=120)
        p = c.paragraphs[0]
        run = p.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    t_table.rows[1].cells[0].text = "{%tr for t in timeline_events %}"
    tr2 = t_table.rows[2]
    tr2.cells[0].width = Inches(2.2)
    tr2.cells[1].width = Inches(1.8)
    tr2.cells[2].width = Inches(2.5)
    tr2.cells[0].text = "{{ t.timestamp }}"
    tr2.cells[1].text = "{{ t.event_type }}"
    tr2.cells[2].text = "{{ t.description }}"
    for c in tr2.cells:
        set_cell_margins(c)
    t_table.rows[3].cells[0].text = "{%tr endfor %}"

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


DEFAULT_HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>Orochi Report - {{ case.name }}</title>
  <link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/@fortawesome/fontawesome-free@6.7.2/css/all.min.css">
  <style>
    @media print {
      body { background-color: #ffffff !important; color: #111827 !important; padding: 0 !important; }
      .no-print { display: none !important; }
      .page-break { page-break-before: always; }
      .shadow-sm, .shadow-md, .shadow-lg { box-shadow: none !important; }
    }
    body {
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
      line-height: 1.5;
      background-color: #f9fafb;
      color: #111827;
      margin: 0;
      padding: 2rem 1rem;
    }
    .container {
      max-width: 900px;
      margin: 0 auto;
      background: #ffffff;
      border: 1px solid #e5e7eb;
      border-radius: 0.75rem;
      padding: 2.5rem;
      box-shadow: 0 1px 3px 0 rgba(0, 0, 0, 0.1);
    }
    .header {
      border-bottom: 2px solid #2563eb;
      padding-bottom: 1.5rem;
      margin-bottom: 2rem;
      display: flex;
      justify-content: space-between;
      align-items: flex-start;
    }
    .header h1 {
      margin: 0;
      font-size: 1.75rem;
      font-weight: 800;
      color: #1e3a8a;
      letter-spacing: -0.025em;
    }
    .header .subtitle {
      color: #4b5563;
      font-size: 1.05rem;
      margin-top: 0.35rem;
      font-weight: 500;
    }
    .badge {
      display: inline-block;
      padding: 0.25rem 0.65rem;
      border-radius: 9999px;
      font-size: 0.75rem;
      font-weight: 700;
      text-transform: uppercase;
      letter-spacing: 0.05em;
    }
    .badge-critical { background: #fee2e2; color: #991b1b; }
    .badge-high { background: #ffedd5; color: #9a3412; }
    .badge-medium { background: #fef3c7; color: #92400e; }
    .badge-low { background: #dbeafe; color: #1e40af; }
    .badge-info { background: #f3f4f6; color: #374151; }
    .meta-grid {
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
      gap: 1rem;
      background: #f8fafc;
      padding: 1.25rem;
      border-radius: 0.5rem;
      border: 1px solid #e2e8f0;
      margin-bottom: 2rem;
    }
    .meta-item .meta-label {
      font-size: 0.75rem;
      font-weight: 600;
      color: #64748b;
      text-transform: uppercase;
    }
    .meta-item .meta-value {
      font-size: 0.95rem;
      font-weight: 600;
      color: #0f172a;
      margin-top: 0.15rem;
    }
    h2 {
      font-size: 1.25rem;
      font-weight: 700;
      color: #1e293b;
      border-bottom: 1px solid #e2e8f0;
      padding-bottom: 0.5rem;
      margin-top: 2.25rem;
      margin-bottom: 1rem;
      display: flex;
      align-items: center;
      gap: 0.5rem;
    }
    table {
      width: 100%;
      border-collapse: collapse;
      margin-bottom: 1.5rem;
      font-size: 0.875rem;
    }
    th {
      background-color: #1e3a8a;
      color: #ffffff;
      text-align: left;
      padding: 0.65rem 0.85rem;
      font-weight: 600;
    }
    td {
      padding: 0.65rem 0.85rem;
      border-bottom: 1px solid #e2e8f0;
      color: #334155;
    }
    tr:nth-child(even) td {
      background-color: #f8fafc;
    }
    .ai-box {
      background: linear-gradient(to right, #eff6ff, #f0fdf4);
      border: 1px solid #bfdbfe;
      border-radius: 0.5rem;
      padding: 1.25rem 1.5rem;
      margin-bottom: 2rem;
      color: #1e293b;
    }
    .ai-box h3 {
      margin-top: 0;
      font-size: 1rem;
      font-weight: 700;
      color: #1d4ed8;
      display: flex;
      align-items: center;
      gap: 0.5rem;
    }
    .btn-print {
      background: #2563eb;
      color: white;
      border: none;
      padding: 0.5rem 1rem;
      border-radius: 0.375rem;
      font-weight: 600;
      cursor: pointer;
      display: inline-flex;
      align-items: center;
      gap: 0.5rem;
    }
    .btn-print:hover { background: #1d4ed8; }
  </style>
</head>
<body>
  <div class="container">
    <div class="header">
      <div>
        <h1>OROCHI INCIDENT REPORT</h1>
        <div class="subtitle">{{ case.name }}</div>
      </div>
      <div class="no-print">
        <button class="btn-print" onclick="window.print()">
          <i class="fa-solid fa-print"></i> Print / Save PDF
        </button>
      </div>
    </div>

    <div class="meta-grid">
      <div class="meta-item">
        <div class="meta-label">Status</div>
        <div class="meta-value">{{ case.status }}</div>
      </div>
      <div class="meta-item">
        <div class="meta-label">Lead Investigator</div>
        <div class="meta-value">{{ case.user.username }}</div>
      </div>
      <div class="meta-item">
        <div class="meta-label">Created At</div>
        <div class="meta-value">{{ case.created_at|date:"Y-m-d H:i" }}</div>
      </div>
      <div class="meta-item">
        <div class="meta-label">Total Findings</div>
        <div class="meta-value">{{ findings|length }}</div>
      </div>
    </div>

    {% if case.description %}
    <h2><i class="fa-solid fa-align-left text-blue-600"></i> Case Description</h2>
    <p>{{ case.description }}</p>
    {% endif %}

    {% if ai_summary %}
    <div class="ai-box">
      <h3><i class="fa-solid fa-wand-magic-sparkles"></i> AI Executive Summary</h3>
      {% if ai_summary_html %}
        {{ ai_summary_html|safe }}
      {% else %}
        <p style="white-space: pre-wrap;">{{ ai_summary }}</p>
      {% endif %}
    </div>
    {% endif %}

    <h2><i class="fa-solid fa-microchip text-blue-600"></i> Evidence Attached ({{ evidences|length }})</h2>
    {% if evidences %}
    <table>
      <thead>
        <tr>
          <th>Evidence Name</th>
          <th>Memory Dump</th>
          <th>Description</th>
        </tr>
      </thead>
      <tbody>
        {% for e in evidences %}
        <tr>
          <td style="font-weight: 600;">{{ e.name }}</td>
          <td>{{ e.dump.name|default:"-" }}</td>
          <td>{{ e.description|default:"-" }}</td>
        </tr>
        {% endfor %}
      </tbody>
    </table>
    {% else %}
    <p style="color: #64748b;">No evidence items attached to this case.</p>
    {% endif %}

    <h2><i class="fa-solid fa-shield-virus text-blue-600"></i> Forensic Findings ({{ findings|length }})</h2>
    {% if findings %}
    <table>
      <thead>
        <tr>
          <th style="width: 120px;">Severity</th>
          <th style="width: 160px;">MITRE ATT&CK</th>
          <th>Note / Evidence</th>
        </tr>
      </thead>
      <tbody>
        {% for f in findings %}
        <tr>
          <td>
            <span class="badge badge-{{ f.severity|lower }}">{{ f.severity }}</span>
          </td>
          <td style="font-family: monospace; font-weight: 600;">{{ f.mitre_attack_technique|default:"-" }}</td>
          <td>{{ f.note }}</td>
        </tr>
        {% endfor %}
      </tbody>
    </table>
    {% else %}
    <p style="color: #64748b;">No findings recorded for this case.</p>
    {% endif %}

    <h2><i class="fa-solid fa-clock-rotate-left text-blue-600"></i> Unified Incident Timeline ({{ timeline_events|length }})</h2>
    {% if timeline_events %}
    <table>
      <thead>
        <tr>
          <th style="width: 180px;">Timestamp</th>
          <th style="width: 180px;">Event Type</th>
          <th>Description</th>
        </tr>
      </thead>
      <tbody>
        {% for t in timeline_events %}
        <tr>
          <td style="font-family: monospace; font-size: 0.8rem;">{{ t.timestamp|date:"Y-m-d H:i:s" }}</td>
          <td style="font-weight: 600;">{{ t.event_type }}</td>
          <td>{{ t.description }}</td>
        </tr>
        {% endfor %}
      </tbody>
    </table>
    {% else %}
    <p style="color: #64748b;">No timeline events recorded.</p>
    {% endif %}
  </div>
</body>
</html>
"""


class Command(BaseCommand):
    help = "Seed default report templates (both DOCX and HTML) in Orochi"

    def handle(self, *args, **options):
        # 1. DOCX Template
        docx_bytes = create_default_docx_bytes()
        docx_tpl, created = ReportTemplate.objects.get_or_create(
            name="Default Word Report (.docx)",
            defaults={"description": "Default professional Word document template for incident reports."},
        )
        docx_tpl.template.save("default_report.docx", ContentFile(docx_bytes), save=True)
        docx_tpl.save()
        self.stdout.write(self.style.SUCCESS(f"Seeded DOCX template: {docx_tpl.name} (pk={docx_tpl.pk})"))

        # Also update any existing 'Default' template to use this docx if it exists
        default_tpl = ReportTemplate.objects.filter(name="Default").first()
        if default_tpl:
            default_tpl.template.save("default.docx", ContentFile(docx_bytes), save=True)
            default_tpl.description = "Default professional Word document template."
            default_tpl.save()
            self.stdout.write(self.style.SUCCESS(f"Updated existing template: Default (pk={default_tpl.pk})"))

        # 2. HTML Template
        html_tpl, created = ReportTemplate.objects.get_or_create(
            name="Default HTML Report (.html)",
            defaults={"description": "Default responsive HTML incident report template."},
        )
        html_tpl.template.save("default_report.html", ContentFile(DEFAULT_HTML_TEMPLATE.encode("utf-8")), save=True)
        html_tpl.save()
        self.stdout.write(self.style.SUCCESS(f"Seeded HTML template: {html_tpl.name} (pk={html_tpl.pk})"))
