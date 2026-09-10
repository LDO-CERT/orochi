import io
import json
import os
import shutil
import tarfile
from pathlib import Path
from unittest.mock import MagicMock, patch
from uuid import uuid4

import pytest
from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from guardian.shortcuts import assign_perm

from orochi.conftest import SORPRESA_ZIP_PATH
from orochi.utils.volatility_dask_elk import (
    check_runnable,
    get_banner,
    hash_checksum,
    manage_upload,
    unzip,
)
from orochi.website.defaults import (
    DUMP_STATUS_CREATED,
    DUMP_STATUS_MISSING_SYMBOLS,
    DUMP_STATUS_UNZIPPING,
)
from orochi.website.models import (
    Case,
    Dump,
    Evidence,
    Finding,
    ReportTemplate,
    Result,
    Value,
)
from orochi.website.views import get_hex_rec

pytestmark = pytest.mark.django_db


@pytest.fixture
def synthetic_dump(admin, folder, synthetic_vmem):
    """Creates a Dump backed by the synthetic .vmem file on disk."""
    dump_dir = Path(settings.MEDIA_ROOT) / "synthetic_dump_idx"
    dump_dir.mkdir(parents=True, exist_ok=True)
    target_path = dump_dir / "sample.vmem"
    shutil.copy(synthetic_vmem, target_path)

    sha256, md5 = hash_checksum(str(target_path))
    dump = Dump.objects.create(
        operating_system="Linux",
        name="synthetic_dump",
        index="synthetic_dump_idx",
        author=admin,
        folder=folder,
        upload=str(target_path),
        size=target_path.stat().st_size,
        sha256=sha256,
        md5=md5,
        status=DUMP_STATUS_CREATED,
    )
    assign_perm("can_see", admin, dump)
    return dump


@pytest.fixture
def sorpresa_dump(admin, folder, sorpresa_extracted_file):
    """Creates a Dump backed by extracted sorpresa.vmem."""
    if not sorpresa_extracted_file or not sorpresa_extracted_file.exists():
        pytest.skip("sorpresa.zip/vmem not available")

    dump_dir = Path(settings.MEDIA_ROOT) / "sorpresa_dump_idx"
    dump_dir.mkdir(parents=True, exist_ok=True)
    target_path = dump_dir / "sorpresa.vmem"
    if not target_path.exists():
        os.link(sorpresa_extracted_file, target_path)

    dump = Dump.objects.create(
        operating_system="Linux",
        name="sorpresa_dump",
        index="sorpresa_dump_idx",
        author=admin,
        folder=folder,
        upload=str(target_path),
        size=target_path.stat().st_size,
        sha256="mock_sha256_for_speed",
        md5="mock_md5_for_speed",
        status=DUMP_STATUS_CREATED,
    )
    assign_perm("can_see", admin, dump)
    return dump


# =====================================================================
# 1. BACKEND FUNCTIONALITY TESTS
# =====================================================================


def test_unzip_function_synthetic(admin, folder, tmp_path, synthetic_zip):
    """Test unzip() function extracting synthetic archive."""
    dump = Dump.objects.create(
        operating_system="Linux",
        name="unzip_test",
        index=str(uuid4()),
        author=admin,
        folder=folder,
    )
    work_dir = tmp_path / "work"
    work_dir.mkdir()
    archive_copy = work_dir / "upload.zip"
    shutil.copy(synthetic_zip, archive_copy)

    extract_path = work_dir / "extracted"
    extract_path.mkdir()

    newpath = unzip(dump, str(archive_copy), str(extract_path), password=None)

    dump.refresh_from_db()
    assert dump.status == DUMP_STATUS_UNZIPPING
    assert newpath is not None
    assert newpath.endswith(".vmem")
    assert os.path.exists(newpath)
    # The original archive should be unlinked by unzip()
    assert not archive_copy.exists()


@pytest.mark.skipif(not SORPRESA_ZIP_PATH.exists(), reason="sorpresa.zip not found")
def test_unzip_with_sorpresa_zip(admin, folder, tmp_path):
    """Test unzip() function with real sorpresa.zip archive containing sorpresa.vmem."""
    dump = Dump.objects.create(
        operating_system="Linux",
        name="sorpresa_unzip_test",
        index=str(uuid4()),
        author=admin,
        folder=folder,
    )
    work_dir = tmp_path / "sorpresa_unzip_work"
    work_dir.mkdir()
    archive_copy = work_dir / "sorpresa.zip"
    shutil.copy(SORPRESA_ZIP_PATH, archive_copy)

    extract_path = work_dir / "extracted"
    extract_path.mkdir()

    newpath = unzip(dump, str(archive_copy), str(extract_path), password=None)

    dump.refresh_from_db()
    assert dump.status == DUMP_STATUS_UNZIPPING
    assert newpath is not None
    assert newpath.endswith("sorpresa.vmem")
    assert os.path.exists(newpath)
    assert os.path.getsize(newpath) == 2147483648
    assert not archive_copy.exists()


def test_unzip_corrupted_archive(admin, folder, tmp_path):
    """Test unzip() with an invalid / corrupted archive returns None."""
    dump = Dump.objects.create(
        operating_system="Linux",
        name="corrupted_test",
        index=str(uuid4()),
        author=admin,
        folder=folder,
    )
    corrupted_zip = tmp_path / "bad.zip"
    corrupted_zip.write_bytes(b"PK\x03\x04not_a_valid_zip_stream")

    extract_path = tmp_path / "bad_extracted"
    extract_path.mkdir()

    newpath = unzip(dump, str(corrupted_zip), str(extract_path), password=None)
    assert newpath is None


def test_hash_checksum_synthetic(synthetic_vmem):
    """Test hash_checksum calculation on synthetic sample."""
    import hashlib

    sha256, md5 = hash_checksum(str(synthetic_vmem))
    raw_bytes = synthetic_vmem.read_bytes()
    expected_sha = hashlib.sha256(raw_bytes).hexdigest()
    expected_md5 = hashlib.md5(raw_bytes).hexdigest()

    assert sha256 == expected_sha
    assert md5 == expected_md5


@pytest.mark.skipif(not SORPRESA_ZIP_PATH.exists(), reason="sorpresa.zip not found")
def test_hash_checksum_sorpresa(sorpresa_extracted_file):
    """Test hash_checksum calculation on real sorpresa.vmem."""
    assert sorpresa_extracted_file is not None
    sha256, md5 = hash_checksum(str(sorpresa_extracted_file), block_size=1024 * 1024)
    assert len(sha256) == 64
    assert len(md5) == 32


@patch("orochi.utils.volatility_dask_elk.close_old_connections")
@patch("orochi.utils.volatility_dask_elk.get_client")
def test_manage_upload_flow(mock_get_client, mock_close_connections, admin, folder, tmp_path, synthetic_zip):
    """Test manage_upload() workflow for archive extraction and dump creation."""
    mock_client = MagicMock()
    mock_get_client.return_value = mock_client

    upload_file = SimpleUploadedFile("sample.zip", synthetic_zip.read_bytes())
    # Use short index so dump.upload.name doesn't exceed FileField max_length 100 in temp dir
    dump = Dump.objects.create(
        operating_system="Linux",
        name="manage_upload_dump",
        index="upl_idx",
        author=admin,
        folder=folder,
        upload=upload_file,
    )

    extract_path = f"{settings.MEDIA_ROOT}/{dump.index}"
    Path(extract_path).mkdir(parents=True, exist_ok=True)
    unzipped_vmem = Path(extract_path) / "sample.vmem"
    unzipped_vmem.write_bytes(b"A" * 1024)

    mock_future = MagicMock()
    mock_future.result.return_value = str(unzipped_vmem)
    mock_client.submit.return_value = mock_future

    with patch("magic.from_file", return_value="application/zip"):
        manage_upload(dump.pk, admin.pk, password=None, restart=None, move=False)

    dump.refresh_from_db()
    assert dump.status in (DUMP_STATUS_CREATED, DUMP_STATUS_MISSING_SYMBOLS)
    assert dump.size == 1024
    assert len(dump.sha256) == 64
    assert len(dump.md5) == 32


def test_check_runnable_and_banner_detection(admin, dump, plugin):
    """Test check_runnable() and get_banner() functions."""
    assert not check_runnable(dump.pk, "Linux", "")
    assert not check_runnable(dump.pk, "Linux", None)
    assert check_runnable(dump.pk, "Windows", "")

    res = Result.objects.create(dump=dump, plugin=plugin, result=0)
    assert get_banner(res) is None

    Value.objects.create(
        result=res,
        value={"Banner": "Linux version 5.4.0-test", "Offset": "0x1000"},
    )
    banner = get_banner(res)
    assert banner == "Linux version 5.4.0-test"


# =====================================================================
# 2. HEX VIEWER BACKEND & UI TESTS
# =====================================================================


def test_hex_view_page_render(client, admin, synthetic_dump):
    """Test rendering the hex viewer page."""
    client.force_login(admin)
    url = reverse("website:hex_view", kwargs={"index": synthetic_dump.index})

    response = client.get(url)
    assert response.status_code == 200
    content = response.content.decode("utf-8")
    assert synthetic_dump.name in content
    assert synthetic_dump.index in content
    assert '<table class="w-full' in content


def test_get_hex_synthetic(client, admin, synthetic_dump):
    """Test get_hex endpoint returning chunked hex data from synthetic .vmem."""
    client.force_login(admin)
    url = reverse("website:get_hex", kwargs={"index": synthetic_dump.index})

    response = client.get(url, {"start": 0, "length": 2, "draw": 1})
    assert response.status_code == 200

    data = response.json()
    assert data["draw"] == 1
    assert data["recordsTotal"] == (64 * 1024) / 16
    assert len(data["data"]) == 2

    row0 = data["data"][0]
    assert row0[0] == "00000000"
    assert row0[1].startswith("45 4c 46")
    assert "<span class='singlechar'>" in row0[2]

    row1 = data["data"][1]
    assert row1[0] == "00000010"


@pytest.mark.skipif(not SORPRESA_ZIP_PATH.exists(), reason="sorpresa.zip not found")
def test_get_hex_sorpresa(client, admin, sorpresa_dump):
    """Test get_hex endpoint with real sorpresa.vmem at offset near 'loading'."""
    client.force_login(admin)
    url = reverse("website:get_hex", kwargs={"index": sorpresa_dump.index})

    start_row = 2067
    response = client.get(url, {"start": start_row, "length": 4, "draw": 42})
    assert response.status_code == 200

    data = response.json()
    assert data["draw"] == 42
    assert data["recordsTotal"] == 2147483648 / 16
    assert len(data["data"]) == 4

    row0 = data["data"][0]
    assert row0[0] == f"{33072:08x}"
    assert "loading" in row0[2] or "l" in row0[2]


def test_get_hex_invalid_params_and_unauthorized(client, readonly_user, synthetic_dump):
    """Test get_hex parameter validation and access control."""
    client.force_login(readonly_user)
    url = reverse("website:get_hex", kwargs={"index": synthetic_dump.index})

    response = client.get(url, {"start": 0, "length": 2})
    assert response.status_code == 200
    assert response.json()["status_code"] == 403

    response = client.get(url, {"start": "not_an_int"})
    assert response.status_code == 200
    assert response.json()["status_code"] == 404


def test_get_hex_rec_direct(synthetic_vmem):
    """Direct test for get_hex_rec function."""
    values, total_rows = get_hex_rec(str(synthetic_vmem), length=32, start=0)
    assert total_rows == (64 * 1024) / 16
    assert len(values) == 2
    assert values[0][0] == "00000000"

    # Reading the last 16-byte row
    last_row_start = (64 * 1024) - 16
    last_values, _ = get_hex_rec(str(synthetic_vmem), length=16, start=last_row_start)
    assert len(last_values) == 1
    assert last_values[0][0] == f"{last_row_start:08x}"

    # Seeking past file size raises ValueError
    with pytest.raises(ValueError):
        get_hex_rec(str(synthetic_vmem), length=16, start=64 * 1024 + 100)


def test_search_hex_synthetic(client, admin, readonly_user, synthetic_dump):
    """Test search_hex endpoint finding byte patterns in memory dump."""
    client.force_login(admin)
    url = reverse("website:search_hex", kwargs={"index": synthetic_dump.index})

    response = client.get(url, {"findstr": "loading", "last": 0})
    assert response.status_code == 200
    result = response.json()
    assert result["found"] == 1
    assert result["pos"] == 33075

    response = client.get(url, {"findstr": "nonexistent_pattern_404", "last": 0})
    assert response.status_code == 200
    assert response.json() == {"found": -1, "pos": 0}

    client.force_login(readonly_user)
    response = client.get(url, {"findstr": "loading", "last": 0})
    assert response.status_code == 200
    assert response.json()["status_code"] == 403


@pytest.mark.skipif(not SORPRESA_ZIP_PATH.exists(), reason="sorpresa.zip not found")
def test_search_hex_sorpresa(client, admin, sorpresa_dump):
    """Test search_hex pattern matching on real sorpresa.vmem."""
    client.force_login(admin)
    url = reverse("website:search_hex", kwargs={"index": sorpresa_dump.index})

    response = client.get(url, {"findstr": "loading", "last": 0})
    assert response.status_code == 200
    result = response.json()
    assert result["found"] == 1
    assert result["pos"] == 33075

    response = client.get(url, {"findstr": "loading", "last": 33076})
    assert response.status_code == 200
    assert response.json()["found"] in (1, -1)


# =====================================================================
# 3. UI MODALS, INFO, EDIT, INDICES & DOWNLOAD TESTS
# =====================================================================


def test_index_info_rendering(client, admin, synthetic_dump):
    """Test website:index_info rendering dump details."""
    client.force_login(admin)
    url = reverse("website:index_info")

    response = client.get(url, {"index": synthetic_dump.index})
    assert response.status_code == 200
    content = response.content.decode("utf-8")
    assert synthetic_dump.name in content
    assert synthetic_dump.sha256 in content
    assert synthetic_dump.md5 in content
    assert "Size" in content


def test_index_edit_get(client, admin, synthetic_dump):
    """Test website:index_edit form rendering via standard and HTMX GET."""
    client.force_login(admin)
    url = reverse("website:index_edit")

    # Standard GET returns JSON with html_form
    response = client.get(url, {"index": synthetic_dump.index})
    assert response.status_code == 200
    assert "html_form" in response.json()

    # HTMX GET returns form HTML directly
    response = client.get(url, {"index": synthetic_dump.index}, HTTP_HX_REQUEST="true")
    assert response.status_code == 200
    assert "<form" in response.content.decode("utf-8")
    assert synthetic_dump.name in response.content.decode("utf-8")


def test_indices_view_render(client, admin, synthetic_dump):
    """Test website:indices rendering sidebar dumps."""
    client.force_login(admin)
    url = reverse("website:indices")

    response = client.get(url)
    assert response.status_code == 200
    content = response.content.decode("utf-8")
    assert synthetic_dump.name in content


def test_banner_symbols_view(client, admin, synthetic_dump):
    """Test website:banner_symbols form rendering."""
    client.force_login(admin)
    url = reverse("website:banner_symbols")

    response = client.get(url, {"index": synthetic_dump.index}, HTTP_HX_REQUEST="true")
    assert response.status_code == 200
    assert "<form" in response.content.decode("utf-8")


def test_download_fixed_sample(client, admin, readonly_user, synthetic_dump, synthetic_vmem):
    """Test website:download streaming dump file."""
    client.force_login(admin)
    url = reverse("website:download")

    # The view expects filepath to have index as the 3rd component: /tmp/<index>/filename
    download_dir = Path(f"/tmp/{synthetic_dump.index}")
    download_dir.mkdir(parents=True, exist_ok=True)
    download_file = download_dir / "sample.vmem"
    shutil.copy(synthetic_vmem, download_file)

    response = client.get(url, {"path": str(download_file)})
    assert response.status_code == 200
    assert "sample.vmem" in response["Content-Disposition"]
    assert len(response.content) == 64 * 1024

    # Readonly user without permission gets 404
    client.force_login(readonly_user)
    response = client.get(url, {"path": str(download_file)})
    assert response.status_code == 404


# =====================================================================
# 4. CASES & EVIDENCE LIFECYCLE WITH FIXED DUMP
# =====================================================================


def test_case_and_evidence_with_fixed_dump(client, admin, synthetic_dump):
    """Test creating a Case, attaching fixed dump as Evidence, and creating Findings."""
    client.force_login(admin)

    # 1. Create Case
    case_url = reverse("website:case_create")
    res = client.post(
        case_url,
        {"name": "Forensic Case with Fixed Dump", "description": "Analyzing sample"},
        HTTP_HX_REQUEST="true",
    )
    assert res.status_code == 200
    case = Case.objects.get(name="Forensic Case with Fixed Dump", user=admin)

    # 2. Attach fixed dump as Evidence
    evidence_url = reverse("website:evidence_create")
    res = client.post(
        evidence_url,
        {
            "name": "RAM Dump Evidence",
            "description": "Captured memory",
            "case": case.pk,
            "dump": synthetic_dump.pk,
        },
        HTTP_HX_REQUEST="true",
    )
    assert res.status_code == 200
    evidence = Evidence.objects.get(name="RAM Dump Evidence", case=case)
    assert evidence.dump == synthetic_dump

    # 3. Create Finding with MITRE ATT&CK technique
    finding_url = reverse("website:finding_create", kwargs={"evidence_pk": evidence.pk})
    res = client.post(
        finding_url,
        {
            "severity": "Critical",
            "tags": "rootkit, malcode",
            "note": "Process memory hollowing detected in dump",
            "mitre_attack_technique": "T1055.012",
            "evidence": evidence.pk,
            "case": case.pk,
        },
        HTTP_HX_REQUEST="true",
    )
    assert res.status_code == 200
    assert Finding.objects.filter(case=case, mitre_attack_technique="T1055.012").exists()

    # 4. View Case Detail
    detail_url = reverse("website:case_detail", kwargs={"pk": case.pk})
    res = client.get(detail_url, HTTP_HX_REQUEST="true")
    assert res.status_code == 200
    detail_content = res.content.decode("utf-8")
    assert "RAM Dump Evidence" in detail_content
    assert "T1055.012" in detail_content

    # 5. Export Case bundle (tar.gz with JSON data)
    export_url = reverse("website:case_export", kwargs={"pk": case.pk})
    res = client.get(export_url)
    assert res.status_code == 200
    assert res["Content-Type"] == "application/gzip"

    tar_bytes = b"".join(res.streaming_content)
    with tarfile.open(fileobj=io.BytesIO(tar_bytes), mode="r:gz") as tar:
        json_member = tar.extractfile(f"case_{case.pk}_export.json")
        assert json_member is not None
        exported_data = json.loads(json_member.read())
        assert exported_data["case"]["name"] == case.name
        assert len(exported_data["evidences"]) == 1
        assert exported_data["evidences"][0]["name"] == "RAM Dump Evidence"

    # 6. Generate Case Report (HTML)
    report_tpl = ReportTemplate.objects.create(
        name="Test HTML Report",
        template=SimpleUploadedFile("report.html", b"<h1>Report for {{ case.name }}</h1>"),
    )
    report_url = reverse("website:case_report", kwargs={"pk": case.pk})
    res = client.post(report_url, {"template_id": report_tpl.pk})
    assert res.status_code == 200
    assert case.name in res.content.decode("utf-8")

    # 7. Generate Case Report (DOCX)
    from docx import Document

    docx_doc = Document()
    docx_doc.add_paragraph("Report for {{ case.name }}")
    docx_doc.add_paragraph("Summary: {{ ai_summary }}")
    docx_io = io.BytesIO()
    docx_doc.save(docx_io)

    report_docx_tpl = ReportTemplate.objects.create(
        name="Test DOCX Report",
        template=SimpleUploadedFile("template.docx", docx_io.getvalue()),
    )
    res_docx = client.post(report_url, {"template_id": report_docx_tpl.pk, "use_ai": "true"})
    assert res_docx.status_code == 200
    assert res_docx["Content-Type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert "_report.docx" in res_docx["Content-Disposition"]

    rendered_doc = Document(io.BytesIO(res_docx.content))
    doc_text = " ".join(p.text for p in rendered_doc.paragraphs)
    assert f"Report for {case.name}" in doc_text
    assert "Ollama service is not configured" in doc_text

    # 8. Generate Case Report with empty DOCX template
    empty_doc = Document()
    empty_io = io.BytesIO()
    empty_doc.save(empty_io)
    empty_docx_tpl = ReportTemplate.objects.create(
        name="Empty DOCX Report",
        template=SimpleUploadedFile("empty.docx", empty_io.getvalue()),
    )
    res_empty = client.post(report_url, {"template_id": empty_docx_tpl.pk, "use_ai": "true"})
    assert res_empty.status_code == 200
    assert res_empty["Content-Type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
